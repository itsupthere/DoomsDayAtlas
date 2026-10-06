"""County-level grid, water, growing-season and climate-trend factors for US counties,
plus each county's NOAA Weather Radio transmitters -> data/us-extra.json

Replaces the old statewide estimates with measured data:
  gs  short growing season   NOAA nClimDiv county minimum temperature, 1991-2020 normals
  cl  observed climate trend  NOAA nClimDiv: warming (1996-2025 vs 1951-1980) and drying (PDSI)
  wa  water stress           NOAA nClimDiv rainfall + drought frequency (PDSI) + USGS 2015 withdrawals
  gr  power grid fragility   EIA-861 2024 reliability: outage minutes per customer (SAIDI incl.
                             major events) of the utilities serving each county
  nwr NOAA Weather Radio     transmitter frequency, call sign, site and coverage per county (SAME code)

    python3 build_us_extra.py [--fetch]     needs openpyxl
"""
import csv, io, json, math, re, sys, urllib.request, zipfile
from pathlib import Path

ROOT = Path(__file__).parent
RAW = ROOT / "data" / "raw"
CLIMDIV = "https://www.ncei.noaa.gov/monitoring-content/data/us/climdiv/monthly/current/"
NCDC_ST = "AL AZ AR CA CO CT DE FL GA ID IL IN IA KS KY LA ME MD MA MI MN MS MO MT NE NV NH NJ NM NY NC ND OH OK OR PA RI SC SD TN TX UT VT VA WA WV WI WY".split()
NCDC_ST = {f"{i + 1:02d}": s for i, s in enumerate(NCDC_ST)} | {"50": "AK"}
FIPS_ST = {"AL": "01", "AK": "02", "AZ": "04", "AR": "05", "CA": "06", "CO": "08", "CT": "09", "DE": "10", "DC": "11", "FL": "12", "GA": "13", "HI": "15",
           "ID": "16", "IL": "17", "IN": "18", "IA": "19", "KS": "20", "KY": "21", "LA": "22", "ME": "23", "MD": "24", "MA": "25", "MI": "26", "MN": "27",
           "MS": "28", "MO": "29", "MT": "30", "NE": "31", "NV": "32", "NH": "33", "NJ": "34", "NM": "35", "NY": "36", "NC": "37", "ND": "38", "OH": "39",
           "OK": "40", "OR": "41", "PA": "42", "RI": "44", "SC": "45", "SD": "46", "TN": "47", "TX": "48", "UT": "49", "VT": "50", "VA": "51", "WA": "53",
           "WV": "54", "WI": "55", "WY": "56"}


def fetch():
    listing = urllib.request.urlopen(CLIMDIV).read().decode()
    for el in ("tmincy", "tmpccy", "pcpncy", "pdsicy"):
        name = sorted(set(re.findall(rf"climdiv-{el}-v[\d.]+-\d+", listing)))[-1]
        (RAW / f"climdiv-{el}.txt").write_bytes(urllib.request.urlopen(CLIMDIV + name).read())
    (RAW / "usgs_wateruse_2015.csv").write_bytes(urllib.request.urlopen(
        "https://www.sciencebase.gov/catalog/file/get/5af3311be4b0da30c1b245d8?f=__disk__eb%2F74%2Feb%2Feb74ebb41169c76aaf374990bd5a71cac82604c1").read())
    (RAW / "eia861_2024.zip").write_bytes(urllib.request.urlopen("https://www.eia.gov/electricity/data/eia861/zip/f8612024.zip").read())
    req = urllib.request.Request("https://www.weather.gov/source/nwr/JS/ccl-data.js", headers={"User-Agent": "Mozilla/5.0 DoomsDayAtlas"})
    (RAW / "nwr-ccl-data.js").write_bytes(urllib.request.urlopen(req).read())
    print("fetched")


def climdiv(el):
    """{fips: {year: [12 monthly values or None]}}"""
    out = {}
    for line in (RAW / f"climdiv-{el}.txt").read_text().splitlines():
        code = line[:11]
        st = NCDC_ST.get(code[:2])
        if not st or st not in FIPS_ST:
            continue
        fips = FIPS_ST[st] + code[2:5]
        vals = [float(x) for x in line[11:].split()]
        out.setdefault(fips, {})[int(code[7:11])] = [None if v <= -99 else v for v in vals]
    return out


def mean(xs):
    xs = [x for x in xs if x is not None]
    return sum(xs) / len(xs) if xs else None


def clamp(x, a=0.0, b=10.0):
    return max(a, min(b, x))


def frost_free_days(tmin):
    """days per year whose interpolated normal minimum is above 36 F (a common proxy for the frost-free season)"""
    days = [15, 46, 74, 105, 135, 166, 196, 227, 258, 288, 319, 349]
    n = 0
    for d in range(365):
        for i in range(12):
            j = (i + 1) % 12
            d0, d1 = days[i], days[j] + (365 if j == 0 else 0)
            dd = d if d >= d0 else d + 365
            if d0 <= dd < d1:
                f = (dd - d0) / (d1 - d0)
                n += (tmin[i] + (tmin[j] - tmin[i]) * f) > 36
                break
        else:
            n += (tmin[11] + (tmin[0] - tmin[11]) * 0.5) > 36
    return n


def build():
    import openpyxl
    tmin, tmpc, pcpn, pdsi = (climdiv(e) for e in ("tmincy", "tmpccy", "pcpncy", "pdsicy"))
    cd = None
    area = {}
    for r in json.loads((RAW / "nri_counties.json").read_text()):
        area[r["STCOFIPS"]] = r["AREA"] or 1
    # USGS freshwater withdrawals (Mgal/day)
    wd = {}
    with open(RAW / "usgs_wateruse_2015.csv", newline="", encoding="latin-1") as f:
        rows = list(csv.reader(f))
    hdr = rows[1]
    ix = {h: i for i, h in enumerate(hdr)}
    for r in rows[2:]:
        try:
            wd[r[ix["FIPS"]].zfill(5)] = float(r[ix["TO-WGWFr"]] or 0) + float(r[ix["TO-WSWFr"]] or 0)
        except (ValueError, IndexError):
            pass
    # EIA: county outage minutes from the utilities serving it, weighted by utility size
    z = zipfile.ZipFile(RAW / "eia861_2024.zip")
    rel = {}
    wb = openpyxl.load_workbook(io.BytesIO(z.read("Reliability_2024.xlsx")), read_only=True)
    num = lambda v: v if isinstance(v, (int, float)) else None
    for r in list(wb["Reliability_States"].iter_rows(values_only=True))[3:]:
        saidi = num(r[5]) if num(r[5]) is not None else num(r[17])
        cust = num(r[14]) if num(r[14]) is not None else num(r[23])
        if saidi is not None and cust:
            rel[(r[1], r[3])] = (saidi, cust)
    state_saidi = {r[1]: num(r[3]) or num(r[13]) for r in list(wb["State Totals"].iter_rows(values_only=True))[3:] if r[1]}
    names = {}
    for fips, (_, row) in ((f, (None, r)) for f, r in ((r["STCOFIPS"], r) for r in json.loads((RAW / "nri_counties.json").read_text()))):
        names[(row["STATEABBRV"], re.sub(r"[^a-z]", "", row["COUNTY"].lower()))] = fips
    county_util = {}
    wb2 = openpyxl.load_workbook(io.BytesIO(z.read("Service_Territory_2024.xlsx")), read_only=True)
    for r in list(wb2["Counties_States"].iter_rows(values_only=True))[1:]:
        st, cn = r[4], re.sub(r"[^a-z]", "", str(r[5]).lower().replace("saint", "st"))
        f = names.get((st, cn)) or names.get((st, cn.replace("st", "saint", 1)))
        if f and (r[1], st) in rel:
            county_util.setdefault(f, []).append(rel[(r[1], st)])

    out = {}
    for f in sorted(area):
        st = next((s for s, c in FIPS_ST.items() if c == f[:2]), None)
        v = {}
        tn = tmin.get(f)
        if tn:
            norm = [mean([tn[y][m] for y in range(1991, 2021) if y in tn]) for m in range(12)]
            if None not in norm:
                ffd = frost_free_days(norm)
                v["gs"] = clamp((250 - ffd) / (250 - 90) * 10)
                v["_ffd"] = ffd
        tp, pd_ = tmpc.get(f), pdsi.get(f)
        if tp:
            past = mean([mean(tp[y]) for y in range(1951, 1981) if y in tp])
            now = mean([mean(tp[y]) for y in range(1996, 2026) if y in tp])
            if past is not None and now is not None:
                warm = now - past  # deg F
                dry = 0.0
                if pd_:
                    dry = (mean([mean(pd_[y]) for y in range(1951, 1981) if y in pd_]) or 0) - (mean([mean(pd_[y]) for y in range(1996, 2026) if y in pd_]) or 0)
                v["cl"] = clamp(1 + 3.2 * (warm - 0.5) + 1.2 * max(dry, 0))  # ~0.5 F -> 1, ~3 F -> 9
                v["_warm"] = round(warm, 2)
        pc = pcpn.get(f)
        if pc:
            precip = mean([sum(x for x in pc[y] if x is not None) for y in range(1991, 2021) if y in pc])
            drought = 0
            if pd_:
                months = [m for y in range(1991, 2026) if y in pd_ for m in pd_[y] if m is not None]
                drought = sum(m <= -3 for m in months) / max(len(months), 1)
            # withdrawals as a share of rain falling on the county (Mgal/day vs inches/yr over sq mi)
            rain_mgd = precip * area.get(f, 1) * 17.38 / 365 * 0.3 if precip else None  # 1 in on 1 sq mi = 17.38 Mgal; ~30% becomes usable flow
            ratio = wd.get(f, 0) / rain_mgd if rain_mgd else 0
            s_dry = clamp((45 - (precip or 45)) / 35 * 10)
            s_drought = clamp(drought / 0.25 * 10)
            s_use = clamp(3 + 3 * math.log10(max(ratio, 1e-3) / 0.05)) if ratio > 0 else 0
            v["wa"] = clamp(0.4 * s_dry + 0.35 * s_drought + 0.25 * s_use)
        us = county_util.get(f)
        saidi = (sum(s * c for s, c in us) / sum(c for _, c in us)) if us else state_saidi.get(st)
        if saidi:
            v["gr"] = clamp(10 * (math.log10(saidi) - math.log10(60)) / (math.log10(2500) - math.log10(60)))
            v["_saidi"] = round(saidi)
        out[f] = v

    # NOAA Weather Radio transmitters per county (SAME code = 0 + county FIPS)
    txt = (RAW / "nwr-ccl-data.js").read_text(encoding="latin-1")
    sites = json.loads(re.sub(r"^\s*var\s+cclData\s*=\s*", "", txt).rstrip().rstrip(";"))
    nwr = {}
    for s in sites:
        if s.get("status", "").upper() not in ("NORMAL", ""):
            continue
        for c in s["counties"]:
            f = c["same"][1:]
            nwr.setdefault(f, []).append([s["freq"], s.get("callsign", ""), s["sitename"], (c.get("remarks") or "").strip().title()])
    keys = ["gs", "wa", "gr", "cl"]
    data = {"keys": keys, "c": {f: [(-1 if k not in v else int(round(v[k] * 10))) for k in keys] + [v.get("_ffd", -1), v.get("_warm", 0), v.get("_saidi", -1)]
                                for f, v in out.items()}, "nwr": nwr}
    (ROOT / "data" / "us-extra.json").write_text(json.dumps(data, separators=(",", ":")))
    have = {k: sum(1 for v in out.values() if k in v) for k in keys}
    print("counties with data:", have, "| NWR counties:", len(nwr))
    for f in ("06037", "53033", "38101", "12086", "04013", "48201", "50023", "02020"):
        v = out.get(f, {})
        print(f, {k: round(x, 1) if isinstance(x, float) else x for k, x in v.items()}, [n[:2] for n in nwr.get(f, [])][:3])


if __name__ == "__main__":
    if "--fetch" in sys.argv:
        fetch()
    build()
